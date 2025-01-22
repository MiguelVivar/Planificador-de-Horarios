import tkinter as tk
from tkinter import messagebox
from tkinter import ttk
import openpyxl  # Asegúrate de tener openpyxl instalado
import os

# Función para cargar los datos desde el archivo Excel
def cargar_datos_excel(ruta):
    try:
        # Cargar el archivo Excel
        excel_data = openpyxl.load_workbook(ruta)
        hoja_activa = excel_data.active

        # Leer encabezados
        if hoja_activa.max_row >= 1:
            headers = [cell.value for cell in hoja_activa[1]]  # Encabezados de la primera fila
        else:
            print("El archivo Excel no tiene datos.")
            return [], []

        data = []
        if headers:
            # Leer los datos de las filas
            for row in range(2, hoja_activa.max_row + 1):
                fila = [hoja_activa.cell(row=row, column=col).value for col in range(1, hoja_activa.max_column + 1)]
                data.append(fila)

            return headers, data
        else:
            print("No se pueden obtener los encabezados.")
            return [], []
    except Exception as e:
        print(f"Error al cargar el archivo Excel: {e}")
        return [], []

# Función para guardar los datos en el archivo Excel
def guardar_datos_excel(directorio, curso, ciclo, seccion, dia, empieza, termina, salon, tipo, profesor):
    # Verificar si el directorio existe, si no, crear uno nuevo
    if not os.path.exists(directorio):
        os.makedirs(directorio)

    # Ruta completa del archivo
    ruta_archivo = os.path.join(directorio, "inscripciones.xlsx")

    # Si el archivo no existe, lo creamos
    if not os.path.exists(ruta_archivo):
        wb = openpyxl.Workbook()
        hoja = wb.active
        # Escribir los encabezados
        hoja.append(["Curso", "Ciclo", "Sección", "Día", "Hora de inicio", "Hora de término", "Número de salón", "Tipo de clase", "Profesor"])
    else:
        wb = openpyxl.load_workbook(ruta_archivo)
        hoja = wb.active

    # Agregar los datos a una nueva fila
    hoja.append([curso, ciclo, seccion, dia, empieza, termina, salon, tipo, profesor])

    # Guardar el archivo
    wb.save(ruta_archivo)

# Función para validar los datos de inscripción
def validar_datos(curso, empieza, termina, profesor):
    if not curso or not empieza or not termina or not profesor:
        return "Todos los campos son obligatorios.", False
    if empieza >= termina:
        return "La hora de inicio debe ser menor que la hora de término.", False
    return "Datos validados correctamente.", True

# Función para verificar inscripción
def verificar_inscripcion():
    salon_seleccionado = salon_var.get().split(" - ")[0]  # Obtener solo el nombre del salón sin la capacidad
    for salon in salones_data:
        if salon[0] == salon_seleccionado:  # El primer valor de cada fila corresponde al nombre del salón
            capacidad = salon[1]  # El segundo valor corresponde a la capacidad
            ocupacion = salon[2]  # El tercer valor corresponde a la ocupación actual (supuesto)
            if isinstance(capacidad, (int, float)) and capacidad > 0:
                if ocupacion < capacidad:
                    mensaje = f"Puedes inscribirte en el salón {salon_seleccionado}.\nCapacidad: {capacidad} - Espacios disponibles: {capacidad - ocupacion}."
                else:
                    mensaje = f"El salón {salon_seleccionado} no tiene capacidad disponible."
            else:
                mensaje = f"Capacidad no definida para el salón {salon_seleccionado}."
            inscripcion_label.config(text=mensaje)
            return
    inscripcion_label.config(text="Selecciona un salón para verificar la inscripción.")

# Ruta al archivo Excel
ruta_excel = r"C:\\Users\\Mario\\Desktop\\CLONE-GITHUB\\Planificador-de-Horarios\\data\\datos.xlsx"

# Cargar los datos del archivo Excel
headers, salones_data = cargar_datos_excel(ruta_excel)

# Configuración de la ventana principal
root = tk.Tk()
root.title("Planificador Moderno")
root.geometry("700x800")
root.configure(bg="#2E4053")

# Estilo
style = ttk.Style()
style.theme_use("clam")
style.configure("TButton", font=("Arial", 14, "bold"), padding=10, background="#58D68D", foreground="white")
style.map("TButton", background=[("active", "#45B39D")])
style.configure("TLabel", background="#2E4053", foreground="white", font=("Arial", 12))
style.configure("TEntry", font=("Arial", 12), padding=5)
style.configure("TCombobox", font=("Arial", 12))

# Encabezado
header = tk.Label(root, text="Planificador de Horarios", font=("Arial", 24, "bold"), bg="#2E4053", fg="#F4D03F")
header.pack(pady=20)

# Contenedor principal
frame = ttk.Frame(root, padding=20, style="TFrame")
frame.pack(expand=True, fill="both", padx=20, pady=20)

# Entradas y menús desplegables
elementos = [
    ("Curso:", ttk.Entry(frame)),
    ("Ciclo:", ttk.Combobox(frame, values=["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"], state="readonly")),
    ("Sección:", ttk.Combobox(frame, values=["A", "B"], state="readonly")),
    ("Día:", ttk.Combobox(frame, values=["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"], state="readonly")),
    ("Hora de inicio:", ttk.Combobox(frame, values=["07:45", "08:30", "09:15", "10:00", "10:45", "11:30", "12:15", "13:00"], state="readonly")),
    ("Hora de término:", ttk.Combobox(frame, values=["08:30", "09:15", "10:00", "10:45", "11:30", "12:15", "13:00", "13:45"], state="readonly")),
    ("Número de salón:", ttk.Combobox(frame, values=[f"{salon[0]} - Capacidad: {salon[1]}" for salon in salones_data if salon[0]], state="readonly")),
    ("Tipo de clase:", ttk.Combobox(frame, values=["Teoría", "Práctica"], state="readonly")),
    ("Profesor:", ttk.Entry(frame)),
]

for i, (label_text, widget) in enumerate(elementos):
    ttk.Label(frame, text=label_text).grid(row=i, column=0, sticky="e", padx=10, pady=10)
    widget.grid(row=i, column=1, sticky="w", padx=10, pady=10)

# Asignar entradas específicas a variables
entrada_curso = elementos[0][1]
ciclo_var = elementos[1][1]
seccion_var = elementos[2][1]
dia_var = elementos[3][1]
empieza_var = elementos[4][1]
termina_var = elementos[5][1]
salon_var = elementos[6][1]
tipo_var = elementos[7][1]
entrada_profesor = elementos[8][1]


# Botones de funcionalidad general
def on_validar():
    curso = entrada_curso.get()
    ciclo = ciclo_var.get()
    seccion = seccion_var.get()
    dia = dia_var.get()
    empieza = empieza_var.get()
    termina = termina_var.get()
    salon = salon_var.get()
    tipo = tipo_var.get()
    profesor = entrada_profesor.get()

    # Validar datos
    mensaje, continuar = validar_datos(curso, empieza, termina, profesor)
    if continuar:
        # Guardar los datos en el archivo Excel
        directorio = r"C:\Users\Mario\Desktop\CLONE-GITHUB\Planificador-de-Horarios\data\horarios"  # Ruta que proporcionaste
        guardar_datos_excel(directorio, curso, ciclo, seccion, dia, empieza, termina, salon, tipo, profesor)
        messagebox.showinfo("Éxito", "El horario fue guardado exitosamente.")
    else:
        messagebox.showwarning("Advertencia", mensaje)

boton_guardar = ttk.Button(frame, text="Validar y Guardar", command=on_validar)
boton_guardar.grid(row=len(elementos) + 2, column=0, columnspan=2, pady=20)

def salir():
    if messagebox.askyesno("Salir", "¿Estás seguro de que deseas salir?"):
        root.destroy()

boton_salir = ttk.Button(frame, text="Salir", command=salir)
boton_salir.grid(row=len(elementos) + 3, column=0, columnspan=2, pady=10)

root.mainloop()